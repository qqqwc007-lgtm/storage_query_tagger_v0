from __future__ import annotations

import argparse
import ast
import csv
import json
import math
import re
from collections import Counter, defaultdict
from copy import copy
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.formatting.rule import ColorScaleRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
except Exception:  # pragma: no cover - optional dependency for workbook export
    Workbook = None
    load_workbook = None
    ColorScaleRule = None
    Alignment = Border = Font = PatternFill = Side = None


VISIBLE_SHEETS = [
    "00_Overview",
    "01_Opportunity_Cluster_Matrix",
    "02_Cluster_Drilldown",
    "03_Methodology_And_Definitions",
]

HIDDEN_SHEETS = [
    "cluster_fact_v1",
    "cluster_keyword_bridge_v1",
    "scene_map",
    "need_map",
    "form_map",
    "metric_summary_v1",
]

DEFAULT_WORKBOOK_NAME = "storage_keyword_distribution_dashboard_v1.xlsx"

SCENE_RULES = [
    ("under_sink", ["under sink", "sink organizer", "sink shelf", "under cabinet"]),
    ("closet", ["closet", "wardrobe", "dresser", "garment", "clothes"]),
    ("pantry", ["pantry", "cereal", "spice rack", "snack"]),
    ("bathroom", ["bathroom", "vanity", "toiletry", "cosmetic"]),
    ("kitchen_drawer", ["drawer organizer", "kitchen drawer", "utensil"]),
    ("cabinet", ["cabinet", "cupboard"]),
    ("fridge", ["fridge", "refrigerator", "freezer"]),
    ("desk", ["desk", "office", "paper", "document"]),
    ("garage", ["garage", "tool", "workbench"]),
    ("laundry", ["laundry", "washer", "dryer"]),
    ("entryway", ["entryway", "mudroom", "hallway"]),
    ("bedroom", ["bedroom", "nightstand"]),
    ("travel", ["travel", "suitcase", "luggage"]),
]

NEED_RULES = [
    ("space_saving", ["space saving", "save space", "slim", "narrow", "compact"]),
    ("visibility", ["clear", "transparent", "see through", "visible", "label"]),
    ("accessibility", ["pull out", "sliding", "rotating", "turntable", "lazy susan", "easy access"]),
    ("stackability", ["stackable", "stacking", "tiered"]),
    ("protection", ["dustproof", "airtight", "sealed", "waterproof", "odor"]),
    ("category_sorting", ["sort", "sorting", "separator", "divider", "compartment"]),
    ("portability", ["portable", "travel", "handle", "carry"]),
    ("organization", ["organizer", "organise", "organize", "storage", "tidy"]),
]

FORM_RULES = [
    ("pull_out_drawer", ["pull out drawer", "sliding drawer", "pull out"]),
    ("drawer_organizer", ["drawer organizer", "drawer tray", "drawer insert"]),
    ("bin", ["bin", "storage bin", "container"]),
    ("basket", ["basket", "wire basket", "woven basket"]),
    ("box", ["box", "shoe box", "file box"]),
    ("rack", ["rack", "shelf", "tiered rack", "spice rack"]),
    ("caddy", ["caddy", "shower caddy"]),
    ("bag", ["bag", "vacuum bag", "hanging bag"]),
    ("cart", ["cart", "rolling cart"]),
    ("divider", ["divider", "separator", "partition"]),
    ("hook", ["hook", "hanger", "hanging"]),
    ("jar", ["jar", "canister", "dispenser"]),
    ("organizer", ["organizer", "organiser"]),
]

OBJECT_RULES = [
    ("shoes", ["shoe", "shoes", "sneaker", "boot"]),
    ("clothes", ["clothes", "shirt", "pants", "sweater", "garment"]),
    ("toys", ["toy", "lego", "doll"]),
    ("spices", ["spice", "seasoning"]),
    ("snacks", ["snack", "chip", "cracker"]),
    ("cans", ["can", "soda", "beverage"]),
    ("lids", ["lid", "pot lid"]),
    ("cosmetics", ["cosmetic", "makeup", "skincare"]),
    ("toiletries", ["toiletry", "toothbrush", "toothpaste"]),
    ("documents", ["document", "paper", "file", "mail"]),
    ("bags", ["bag", "handbag", "purse"]),
    ("towels", ["towel", "linen"]),
]

GENERIC_STOPWORDS = {
    "and",
    "for",
    "with",
    "the",
    "a",
    "an",
    "of",
    "to",
    "in",
    "on",
    "set",
    "pack",
    "home",
    "kitchen",
    "bathroom",
    "closet",
    "storage",
    "organizer",
    "organisers",
    "organizer",
    "organizers",
}

SPECIFICITY_FACTORS = {
    "scene": 1.00,
    "object": 1.10,
    "form": 0.95,
    "generic": 0.85,
}

QUALITY_FACTORS = {
    "ready": 1.00,
    "candidate": 0.72,
    "backlog": 0.48,
    "watchlist": 0.28,
}

ROW_OUTCOME_SYNONYMS = {
    "mapped": "mapped",
    "complete": "mapped",
    "ready": "mapped",
    "partial": "partial",
    "partially_mapped": "partial",
    "candidate": "candidate",
    "discovered": "candidate",
    "new_value": "candidate",
    "unmapped": "unmapped",
    "ambiguous": "ambiguous",
    "unclear": "ambiguous",
    "conflict": "ambiguous",
}

KEYWORD_FIELDS = ["keyword", "keyword_text", "kw", "query", "search_term", "term", "title"]
RANK_FIELDS = ["rank", "organic_rank", "search_rank", "keyword_rank", "position"]
SEARCH_VOLUME_FIELDS = ["search_volume", "volume", "estimated_search_volume"]
TOP100K_FIELDS = ["top100k_flag", "is_top100k", "top_100k"]
SCENE_FIELDS = ["scene_rollup_key", "scenario_rollup_key", "scene_key", "scene", "scenario", "use_scene"]
NEED_FIELDS = ["function_key", "need_rollup_key", "need_key", "need", "function_need", "function"]
FORM_FIELDS = ["form_rollup_key", "form_key", "form", "storage_form", "product_form"]
ANCHOR_TYPE_FIELDS = ["anchor_type"]
ANCHOR_VALUE_FIELDS = ["anchor_value"]
OBJECT_FIELDS = ["object_key", "object", "target_object"]
OUTCOME_FIELDS = ["row_outcome", "mapping_status", "status"]
AMBIGUOUS_FIELDS = ["ambiguous", "is_ambiguous"]
MATCHED_ALIAS_FIELDS = ["matched_aliases", "matched_alias", "aliases"]
EVIDENCE_FIELDS = ["evidence", "evidence_summary", "reason"]
UNMAPPED_FIELDS = ["unmapped_phrases", "unmapped_phrase", "unknown_phrases"]


@dataclass(frozen=True)
class DashboardBuildConfig:
    kw_csv_path: Path
    output_dir: Path
    template_path: Path | None = None
    snapshot_date: str | None = None
    marketplace: str = "US"
    taxonomy_version: str = "storage_v1"
    cluster_key_version: str = "v1"
    workbook_name: str = DEFAULT_WORKBOOK_NAME
    write_workbook: bool = True


def _clean_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _slugify(value: Any) -> str:
    text = _clean_text(value).lower()
    text = re.sub(r"[^a-z0-9]+", "_", text)
    return text.strip("_")


def _titleize(key: str) -> str:
    return key.replace("_", " ").title()


def _as_bool(value: Any) -> bool:
    text = _clean_text(value).lower()
    return text in {"1", "true", "yes", "y", "mapped", "partial", "ambiguous"}


def _as_float(value: Any) -> float | None:
    text = _clean_text(value)
    if not text:
        return None
    text = text.replace(",", "")
    try:
        return float(text)
    except ValueError:
        return None


def _as_int(value: Any) -> int | None:
    number = _as_float(value)
    if number is None:
        return None
    return int(number)


def _parse_listish(value: Any) -> list[str]:
    if value is None:
        return []
    if isinstance(value, (list, tuple, set)):
        return [_clean_text(item) for item in value if _clean_text(item)]
    text = _clean_text(value)
    if not text or text.lower() in {"none", "null", "nan", "[]", "{}"}:
        return []
    if text.startswith("[") or text.startswith("{") or text.startswith("("):
        for parser in (json.loads, ast.literal_eval):
            try:
                parsed = parser(text)
            except Exception:
                continue
            if isinstance(parsed, dict):
                parsed = list(parsed.values())
            if isinstance(parsed, (list, tuple, set)):
                return [_clean_text(item) for item in parsed if _clean_text(item)]
    for separator in ("|", ";", "/", ","):
        if separator in text:
            parts = [_clean_text(part) for part in text.split(separator)]
            return [part for part in parts if part]
    return [text]


def _pick_first(row: dict[str, Any], field_names: Iterable[str]) -> str:
    for field_name in field_names:
        value = _clean_text(row.get(field_name))
        if value:
            return value
    return ""


def _percent_rank(values: list[float], current_value: float) -> float:
    if not values:
        return 0.0
    if len(values) == 1:
        return 1.0
    ordered = sorted(values)
    lower = sum(1 for value in ordered if value < current_value)
    equal = sum(1 for value in ordered if value == current_value)
    midpoint = lower + max(equal - 1, 0) / 2.0
    return midpoint / (len(ordered) - 1)


def _read_csv_rows(csv_path: Path) -> list[dict[str, str]]:
    with csv_path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def _write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if not rows:
        path.write_text("", encoding="utf-8")
        return
    fieldnames = list(rows[0].keys())
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        for row in rows:
            writer.writerow(row)


def _extract_dimension(
    row: dict[str, Any],
    field_names: list[str],
    rules: list[tuple[str, list[str]]],
    keyword_text: str,
) -> tuple[list[str], dict[str, str] | None]:
    direct_values: list[str] = []
    source_field = ""
    raw_value = ""
    for field_name in field_names:
        values = _parse_listish(row.get(field_name))
        if values:
            direct_values = [_slugify(value) for value in values if _slugify(value)]
            source_field = field_name
            raw_value = " | ".join(values)
            break

    if direct_values:
        resolved = list(dict.fromkeys(direct_values))
        return resolved, {
            "source_field": source_field,
            "raw_value": raw_value,
            "rollup_key": resolved[0],
            "rollup_label": _titleize(resolved[0]),
            "mapping_method": "direct",
        }

    inferred = [key for key, patterns in rules if any(_matches_pattern(keyword_text, pattern) for pattern in patterns)]
    inferred = list(dict.fromkeys(inferred))
    if inferred:
        return inferred, {
            "source_field": "keyword_inference",
            "raw_value": keyword_text,
            "rollup_key": inferred[0],
            "rollup_label": _titleize(inferred[0]),
            "mapping_method": "inferred",
        }
    return [], None


def _extract_object_key(row: dict[str, Any], keyword_text: str) -> str:
    explicit = _pick_first(row, OBJECT_FIELDS)
    if explicit:
        return _slugify(explicit)
    for object_key, patterns in OBJECT_RULES:
        if any(_matches_pattern(keyword_text, pattern) for pattern in patterns):
            return object_key
    return "generic"


def _matches_pattern(text: str, pattern: str) -> bool:
    normalized_text = f" {_normalize_search_text(text)} "
    normalized_pattern = _normalize_search_text(pattern)
    if not normalized_pattern:
        return False
    return f" {normalized_pattern} " in normalized_text


def _normalize_search_text(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", text.lower()).strip()


def _infer_anchor_type(row: dict[str, Any], scene_keys: list[str], object_key: str, form_keys: list[str]) -> str:
    explicit = _slugify(_pick_first(row, ANCHOR_TYPE_FIELDS))
    if explicit:
        return explicit
    if scene_keys and scene_keys[0] != "unspecified_scene":
        return "scene"
    if object_key and object_key != "generic":
        return "object"
    if form_keys and form_keys[0] != "unspecified_form":
        return "form"
    return "generic"


def _infer_anchor_value(row: dict[str, Any], anchor_type: str, scene_keys: list[str], object_key: str, form_keys: list[str]) -> str:
    explicit = _slugify(_pick_first(row, ANCHOR_VALUE_FIELDS))
    if explicit:
        return explicit
    if anchor_type == "scene":
        return scene_keys[0] if scene_keys else "general"
    if anchor_type == "object":
        return object_key or "generic"
    if anchor_type == "form":
        return form_keys[0] if form_keys else "generic"
    return "general"


def _derive_row_outcome(
    row: dict[str, Any],
    scene_keys: list[str],
    need_keys: list[str],
    form_keys: list[str],
    unmatched_terms: list[str],
) -> str:
    explicit = _slugify(_pick_first(row, OUTCOME_FIELDS))
    if explicit:
        return ROW_OUTCOME_SYNONYMS.get(explicit, explicit)
    if any(_as_bool(row.get(field_name)) for field_name in AMBIGUOUS_FIELDS):
        return "ambiguous"
    dimension_count = int(bool(scene_keys)) + int(bool(need_keys)) + int(bool(form_keys))
    if dimension_count == 3 and not unmatched_terms:
        return "mapped"
    if dimension_count >= 1:
        return "partial"
    return "candidate"


def _dashboard_bucket(row_outcome: str) -> str:
    if row_outcome in {"mapped", "partial"}:
        return "main_pool"
    if row_outcome == "ambiguous":
        return "watchlist"
    return "backlog"


def _build_cluster_key(
    cluster_key_version: str,
    anchor_type: str,
    anchor_value: str,
    object_key: str,
    form_rollup_key: str,
    function_key: str,
) -> str:
    return "|".join(
        [
            cluster_key_version,
            f"anchor_type={anchor_type or 'generic'}",
            f"anchor_value={anchor_value or 'general'}",
            f"object_key={object_key or 'generic'}",
            f"form_rollup_key={form_rollup_key or 'unspecified_form'}",
            f"function_key={function_key or 'unspecified_need'}",
        ]
    )


def _cluster_quality(rows: list[dict[str, Any]], scene_key: str, function_key: str, form_key: str) -> str:
    row_counts = Counter(row["row_outcome"] for row in rows)
    total = len(rows) or 1
    ambiguity_rate = row_counts["ambiguous"] / total
    mapped_like_count = row_counts["mapped"] + row_counts["partial"]
    fully_scoped = all(
        key and not key.startswith("unspecified")
        for key in (scene_key, function_key, form_key)
    )
    if ambiguity_rate >= 0.40:
        return "watchlist"
    object_scoped = any(row["object_key"] != "generic" for row in rows)
    if fully_scoped and object_scoped and row_counts["mapped"] >= 2 and mapped_like_count / total >= 0.75:
        return "ready"
    if fully_scoped and (mapped_like_count >= 2 or row_counts["candidate"] + row_counts["unmapped"] >= 1):
        return "candidate"
    if mapped_like_count >= 1:
        return "backlog"
    return "watchlist" if row_counts["ambiguous"] else "backlog"


def _rank_value(row: dict[str, Any]) -> int:
    rank = row["rank"]
    return rank if isinstance(rank, int) else 10**9


def _summarize_evidence(rows: list[dict[str, Any]]) -> str:
    parts: list[str] = []
    for row in sorted(rows, key=_rank_value)[:5]:
        parts.extend(row["matched_aliases"])
        if row["evidence"]:
            parts.append(row["evidence"])
        parts.extend(row["unmapped_phrases"])
    deduped = []
    seen = set()
    for part in parts:
        cleaned = _clean_text(part)
        if cleaned and cleaned not in seen:
            seen.add(cleaned)
            deduped.append(cleaned)
    summary = " | ".join(deduped[:8])
    return summary[:240]


def _normalized_rows(
    raw_rows: list[dict[str, Any]],
    config: DashboardBuildConfig,
) -> tuple[list[dict[str, Any]], dict[str, list[dict[str, str]]]]:
    scene_map_rows: list[dict[str, str]] = []
    need_map_rows: list[dict[str, str]] = []
    form_map_rows: list[dict[str, str]] = []
    normalized: list[dict[str, Any]] = []

    for index, raw_row in enumerate(raw_rows, start=1):
        keyword = _pick_first(raw_row, KEYWORD_FIELDS)
        keyword = keyword or f"row_{index}"
        snapshot_date = _pick_first(raw_row, ["snapshot_date"]) or config.snapshot_date or date.today().isoformat()
        marketplace = _pick_first(raw_row, ["marketplace", "site"]) or config.marketplace
        taxonomy_version = _pick_first(raw_row, ["taxonomy_version"]) or config.taxonomy_version
        rank = None
        for field_name in RANK_FIELDS:
            rank = _as_int(raw_row.get(field_name))
            if rank is not None:
                break
        top100k_flag = any(_as_bool(raw_row.get(field_name)) for field_name in TOP100K_FIELDS)
        search_volume = None
        for field_name in SEARCH_VOLUME_FIELDS:
            search_volume = _as_float(raw_row.get(field_name))
            if search_volume is not None:
                break

        scene_keys, scene_map_entry = _extract_dimension(raw_row, SCENE_FIELDS, SCENE_RULES, keyword)
        need_keys, need_map_entry = _extract_dimension(raw_row, NEED_FIELDS, NEED_RULES, keyword)
        form_keys, form_map_entry = _extract_dimension(raw_row, FORM_FIELDS, FORM_RULES, keyword)

        scene_key = scene_keys[0] if scene_keys else "unspecified_scene"
        function_key = need_keys[0] if need_keys else "unspecified_need"
        form_rollup_key = form_keys[0] if form_keys else "unspecified_form"
        object_key = _extract_object_key(raw_row, keyword)
        anchor_type = _infer_anchor_type(raw_row, scene_keys, object_key, form_keys)
        anchor_value = _infer_anchor_value(raw_row, anchor_type, scene_keys, object_key, form_keys)
        matched_aliases = _parse_listish(_pick_first(raw_row, MATCHED_ALIAS_FIELDS))
        evidence = _pick_first(raw_row, EVIDENCE_FIELDS)
        unmapped_phrases = _parse_listish(_pick_first(raw_row, UNMAPPED_FIELDS))
        row_outcome = _derive_row_outcome(raw_row, scene_keys, need_keys, form_keys, unmapped_phrases)
        cluster_key = _build_cluster_key(
            config.cluster_key_version,
            anchor_type,
            anchor_value,
            object_key,
            form_rollup_key,
            function_key,
        )

        if scene_map_entry:
            scene_map_rows.append({"dimension": "scene", **scene_map_entry})
        if need_map_entry:
            need_map_rows.append({"dimension": "need", **need_map_entry})
        if form_map_entry:
            form_map_rows.append({"dimension": "form", **form_map_entry})

        normalized.append(
            {
                "source_row_number": index,
                "snapshot_date": snapshot_date,
                "marketplace": marketplace,
                "taxonomy_version": taxonomy_version,
                "cluster_key_version": config.cluster_key_version,
                "keyword": keyword,
                "rank": rank,
                "top100k_flag": top100k_flag,
                "search_volume": search_volume,
                "anchor_type": anchor_type,
                "anchor_value": anchor_value,
                "object_key": object_key,
                "scene_key": scene_key,
                "scene_label": _titleize(scene_key),
                "function_key": function_key,
                "function_label": _titleize(function_key),
                "form_rollup_key": form_rollup_key,
                "form_label": _titleize(form_rollup_key),
                "row_outcome": row_outcome,
                "dashboard_bucket": _dashboard_bucket(row_outcome),
                "matched_aliases": matched_aliases,
                "evidence": evidence,
                "unmapped_phrases": unmapped_phrases,
                "cluster_key": cluster_key,
            }
        )

    return normalized, {
        "scene_map": _dedupe_mapping_rows(scene_map_rows),
        "need_map": _dedupe_mapping_rows(need_map_rows),
        "form_map": _dedupe_mapping_rows(form_map_rows),
    }


def _dedupe_mapping_rows(rows: list[dict[str, str]]) -> list[dict[str, str]]:
    deduped: list[dict[str, str]] = []
    seen: set[tuple[str, ...]] = set()
    for row in rows:
        key = tuple(row.get(column, "") for column in ["dimension", "source_field", "raw_value", "rollup_key", "mapping_method"])
        if key in seen:
            continue
        seen.add(key)
        deduped.append(row)
    return sorted(deduped, key=lambda row: (row["dimension"], row["rollup_key"], row["raw_value"]))


def materialize_dashboard_contracts(
    raw_rows: list[dict[str, Any]],
    config: DashboardBuildConfig,
) -> dict[str, list[dict[str, Any]]]:
    normalized_rows, mapping_tables = _normalized_rows(raw_rows, config)
    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in normalized_rows:
        grouped[row["cluster_key"]].append(row)

    cluster_fact_rows: list[dict[str, Any]] = []
    bridge_rows: list[dict[str, Any]] = []

    for cluster_key, rows in grouped.items():
        rows_sorted = sorted(rows, key=_rank_value)
        representative = rows_sorted[0]
        keyword_count = len(rows)
        row_counts = Counter(row["row_outcome"] for row in rows)
        ranks = [row["rank"] for row in rows if isinstance(row["rank"], int)]
        search_volumes = [row["search_volume"] for row in rows if isinstance(row["search_volume"], (float, int))]
        top100k_keyword_count = sum(1 for row in rows if row["top100k_flag"] or (isinstance(row["rank"], int) and row["rank"] <= 100000))
        best_rank = min(ranks) if ranks else None
        avg_rank = round(sum(ranks) / len(ranks), 2) if ranks else None
        search_volume_sum = round(sum(search_volumes), 2) if search_volumes else 0.0
        search_volume_coverage = round(len(search_volumes) / keyword_count, 4) if keyword_count else 0.0
        mapped_like_share = round((row_counts["mapped"] + row_counts["partial"]) / keyword_count, 4) if keyword_count else 0.0
        top100k_penetration = round(top100k_keyword_count / keyword_count, 4) if keyword_count else 0.0
        rank_weighted_demand = round(sum(1.0 / max(rank, 1) for rank in ranks), 8)
        cluster_quality = _cluster_quality(
            rows,
            representative["scene_key"],
            representative["function_key"],
            representative["form_rollup_key"],
        )
        specificity_factor = SPECIFICITY_FACTORS.get(representative["anchor_type"], SPECIFICITY_FACTORS["generic"])
        representative_keywords = " | ".join(row["keyword"] for row in rows_sorted[:3])
        evidence_summary = _summarize_evidence(rows)
        dashboard_tab = "02_Cluster_Drilldown" if cluster_quality == "watchlist" else "01_Opportunity_Cluster_Matrix"
        cluster_label_parts = [
            representative["scene_label"],
            representative["function_label"],
            representative["form_label"],
        ]
        if representative["object_key"] != "generic":
            cluster_label_parts.append(_titleize(representative["object_key"]))
        cluster_label = " / ".join(part for part in cluster_label_parts if part)

        cluster_fact_row = {
            "snapshot_date": representative["snapshot_date"],
            "marketplace": representative["marketplace"],
            "taxonomy_version": representative["taxonomy_version"],
            "cluster_key_version": representative["cluster_key_version"],
            "cluster_key": cluster_key,
            "cluster_label": cluster_label,
            "anchor_type": representative["anchor_type"],
            "anchor_value": representative["anchor_value"],
            "object_key": representative["object_key"],
            "scene_key": representative["scene_key"],
            "scene_label": representative["scene_label"],
            "function_key": representative["function_key"],
            "function_label": representative["function_label"],
            "form_rollup_key": representative["form_rollup_key"],
            "form_label": representative["form_label"],
            "cluster_quality": cluster_quality,
            "dashboard_tab": dashboard_tab,
            "include_in_top10": False,
            "ready_rank": "",
            "keyword_count": keyword_count,
            "mapped_keyword_count": row_counts["mapped"],
            "partial_keyword_count": row_counts["partial"],
            "ambiguous_keyword_count": row_counts["ambiguous"],
            "candidate_keyword_count": row_counts["candidate"],
            "unmapped_keyword_count": row_counts["unmapped"],
            "mapped_like_share": mapped_like_share,
            "top100k_keyword_count": top100k_keyword_count,
            "top100k_penetration": top100k_penetration,
            "best_rank": best_rank or "",
            "avg_rank": avg_rank or "",
            "search_volume_present_count": len(search_volumes),
            "search_volume_coverage": search_volume_coverage,
            "search_volume_sum": search_volume_sum,
            "rank_weighted_demand": rank_weighted_demand,
            "specificity_factor": specificity_factor,
            "priority_score": 0.0,
            "representative_keyword": representative["keyword"],
            "representative_keywords": representative_keywords,
            "evidence_summary": evidence_summary,
        }
        cluster_fact_rows.append(cluster_fact_row)

        for row in rows_sorted:
            bridge_rows.append(
                {
                    "snapshot_date": row["snapshot_date"],
                    "marketplace": row["marketplace"],
                    "cluster_key": cluster_key,
                    "cluster_label": cluster_label,
                    "keyword": row["keyword"],
                    "rank": row["rank"] or "",
                    "search_volume": row["search_volume"] or "",
                    "row_outcome": row["row_outcome"],
                    "dashboard_bucket": row["dashboard_bucket"],
                    "anchor_type": row["anchor_type"],
                    "anchor_value": row["anchor_value"],
                    "object_key": row["object_key"],
                    "scene_key": row["scene_key"],
                    "function_key": row["function_key"],
                    "form_rollup_key": row["form_rollup_key"],
                    "matched_aliases": " | ".join(row["matched_aliases"]),
                    "evidence": row["evidence"],
                    "unmapped_phrases": " | ".join(row["unmapped_phrases"]),
                    "source_row_number": row["source_row_number"],
                }
            )

    _assign_priority_scores(cluster_fact_rows)
    _mark_top10_ready(cluster_fact_rows)

    metric_summary_rows = _metric_summary_rows(normalized_rows, cluster_fact_rows, config)
    return {
        "cluster_fact_v1": sorted(
            cluster_fact_rows,
            key=lambda row: (_quality_sort_key(row["cluster_quality"]), -float(row["priority_score"]), -int(row["keyword_count"])),
        ),
        "cluster_keyword_bridge_v1": bridge_rows,
        "scene_map": mapping_tables["scene_map"],
        "need_map": mapping_tables["need_map"],
        "form_map": mapping_tables["form_map"],
        "metric_summary_v1": metric_summary_rows,
    }


def _assign_priority_scores(cluster_fact_rows: list[dict[str, Any]]) -> None:
    keyword_counts = [float(row["keyword_count"]) for row in cluster_fact_rows]
    top100k_penetrations = [float(row["top100k_penetration"]) for row in cluster_fact_rows]
    rank_demands = [float(row["rank_weighted_demand"]) for row in cluster_fact_rows]

    for row in cluster_fact_rows:
        keyword_pct = _percent_rank(keyword_counts, float(row["keyword_count"]))
        penetration_pct = _percent_rank(top100k_penetrations, float(row["top100k_penetration"]))
        demand_pct = _percent_rank(rank_demands, float(row["rank_weighted_demand"]))
        raw_score = 0.45 * keyword_pct + 0.35 * penetration_pct + 0.20 * demand_pct
        quality_factor = QUALITY_FACTORS[row["cluster_quality"]]
        specificity_factor = float(row["specificity_factor"])
        row["priority_score"] = round(min(100.0, 100.0 * raw_score * quality_factor * specificity_factor), 2)


def _mark_top10_ready(cluster_fact_rows: list[dict[str, Any]]) -> None:
    ready_rows = sorted(
        [row for row in cluster_fact_rows if row["cluster_quality"] == "ready"],
        key=lambda row: (-float(row["priority_score"]), -int(row["keyword_count"]), row["cluster_label"]),
    )
    for index, row in enumerate(ready_rows[:10], start=1):
        row["include_in_top10"] = True
        row["ready_rank"] = index


def _quality_sort_key(cluster_quality: str) -> int:
    return {"ready": 0, "candidate": 1, "backlog": 2, "watchlist": 3}.get(cluster_quality, 9)


def _metric_summary_rows(
    normalized_rows: list[dict[str, Any]],
    cluster_fact_rows: list[dict[str, Any]],
    config: DashboardBuildConfig,
) -> list[dict[str, Any]]:
    row_counts = Counter(row["row_outcome"] for row in normalized_rows)
    cluster_counts = Counter(row["cluster_quality"] for row in cluster_fact_rows)
    total_rows = len(normalized_rows) or 1
    total_clusters = len(cluster_fact_rows) or 1
    search_volume_coverage = sum(1 for row in normalized_rows if row["search_volume"] is not None) / total_rows
    top100k_share = sum(
        1
        for row in normalized_rows
        if row["top100k_flag"] or (isinstance(row["rank"], int) and row["rank"] <= 100000)
    ) / total_rows
    rows = [
        ("snapshot_date", _resolve_metadata_value(normalized_rows, "snapshot_date", config.snapshot_date or date.today().isoformat()), "Dashboard snapshot date"),
        ("marketplace", _resolve_metadata_value(normalized_rows, "marketplace", config.marketplace), "Marketplace scope"),
        ("total_keywords", len(normalized_rows), "Total keyword rows"),
        ("mapped_keywords", row_counts["mapped"], "Fully mapped keyword rows"),
        ("partial_keywords", row_counts["partial"], "Partially mapped keyword rows"),
        ("ambiguous_keywords", row_counts["ambiguous"], "Ambiguous keyword rows"),
        ("candidate_keywords", row_counts["candidate"], "Candidate keyword rows"),
        ("unmapped_keywords", row_counts["unmapped"], "Explicitly unmapped keyword rows"),
        ("main_pool_keywords", row_counts["mapped"] + row_counts["partial"], "Rows eligible for main pool"),
        ("total_clusters", len(cluster_fact_rows), "Opportunity clusters"),
        ("ready_clusters", cluster_counts["ready"], "Clusters eligible for Top 10"),
        ("candidate_clusters", cluster_counts["candidate"], "Clusters kept for follow-up"),
        ("backlog_clusters", cluster_counts["backlog"], "Clusters needing mapping completion"),
        ("watchlist_clusters", cluster_counts["watchlist"], "Clusters routed to watchlist"),
        ("search_volume_coverage", round(search_volume_coverage, 4), "Coverage of populated search volume"),
        ("top100k_keyword_share", round(top100k_share, 4), "Share of keyword rows ranked within Top100k"),
        ("top10_ready_cluster_share", round(cluster_counts["ready"] / total_clusters, 4), "Share of clusters marked ready"),
        ("generated_at", datetime.now().isoformat(timespec="seconds"), "Artifact build time"),
    ]
    return [
        {
            "metric_name": metric_name,
            "metric_value": metric_value,
            "notes": notes,
        }
        for metric_name, metric_value, notes in rows
    ]


def build_dashboard_from_contracts(
    contract_tables: dict[str, list[dict[str, Any]]],
    config: DashboardBuildConfig,
) -> dict[str, Path]:
    config.output_dir.mkdir(parents=True, exist_ok=True)
    artifact_paths = {
        "cluster_fact_v1": config.output_dir / "cluster_fact_v1.csv",
        "cluster_keyword_bridge_v1": config.output_dir / "cluster_keyword_bridge_v1.csv",
        "scene_map": config.output_dir / "scene_map.csv",
        "need_map": config.output_dir / "need_map.csv",
        "form_map": config.output_dir / "form_map.csv",
        "metric_summary_v1": config.output_dir / "metric_summary_v1.csv",
    }
    for table_name, path in artifact_paths.items():
        _write_csv(path, contract_tables[table_name])

    workbook_path = config.output_dir / config.workbook_name
    if config.write_workbook:
        _write_workbook(
            workbook_path=workbook_path,
            template_path=config.template_path,
            contract_tables=contract_tables,
            source_path=config.kw_csv_path,
            config=config,
        )
    return {**artifact_paths, "workbook": workbook_path}


def build_keyword_dashboard(config: DashboardBuildConfig) -> dict[str, Path]:
    raw_rows = _read_csv_rows(config.kw_csv_path)
    contract_tables = materialize_dashboard_contracts(raw_rows, config)
    return build_dashboard_from_contracts(contract_tables, config)


def _ensure_openpyxl() -> None:
    if Workbook is None or load_workbook is None:
        raise RuntimeError("openpyxl is required to write the workbook artifact. Re-run with --no-workbook or install openpyxl.")


def _write_workbook(
    workbook_path: Path,
    template_path: Path | None,
    contract_tables: dict[str, list[dict[str, Any]]],
    source_path: Path,
    config: DashboardBuildConfig,
) -> None:
    _ensure_openpyxl()
    if template_path and template_path.exists():
        workbook = load_workbook(template_path)
        for sheet_name in list(workbook.sheetnames):
            workbook.remove(workbook[sheet_name])
    else:
        workbook = Workbook()
        workbook.remove(workbook.active)

    overview = workbook.create_sheet(VISIBLE_SHEETS[0])
    matrix_sheet = workbook.create_sheet(VISIBLE_SHEETS[1])
    drilldown_sheet = workbook.create_sheet(VISIBLE_SHEETS[2])
    methodology = workbook.create_sheet(VISIBLE_SHEETS[3])
    hidden_sheets = {name: workbook.create_sheet(name) for name in HIDDEN_SHEETS}
    effective_snapshot_date = _resolve_metadata_value(contract_tables["cluster_fact_v1"], "snapshot_date", config.snapshot_date or date.today().isoformat())
    effective_marketplace = _resolve_metadata_value(contract_tables["cluster_fact_v1"], "marketplace", config.marketplace)
    effective_taxonomy_version = _resolve_metadata_value(contract_tables["cluster_fact_v1"], "taxonomy_version", config.taxonomy_version)

    _populate_overview_sheet(
        overview,
        contract_tables["cluster_fact_v1"],
        contract_tables["metric_summary_v1"],
        config,
        effective_snapshot_date,
        effective_marketplace,
        effective_taxonomy_version,
    )
    _populate_matrix_sheet(matrix_sheet, contract_tables["cluster_fact_v1"])
    _populate_drilldown_sheet(drilldown_sheet, contract_tables["cluster_keyword_bridge_v1"], contract_tables["cluster_fact_v1"])
    _populate_methodology_sheet(methodology, config, source_path, effective_snapshot_date, effective_marketplace, effective_taxonomy_version)
    _populate_contract_sheet(hidden_sheets["cluster_fact_v1"], contract_tables["cluster_fact_v1"])
    _populate_contract_sheet(hidden_sheets["cluster_keyword_bridge_v1"], contract_tables["cluster_keyword_bridge_v1"])
    _populate_contract_sheet(hidden_sheets["scene_map"], contract_tables["scene_map"])
    _populate_contract_sheet(hidden_sheets["need_map"], contract_tables["need_map"])
    _populate_contract_sheet(hidden_sheets["form_map"], contract_tables["form_map"])
    _populate_contract_sheet(hidden_sheets["metric_summary_v1"], contract_tables["metric_summary_v1"])

    for sheet_name in HIDDEN_SHEETS:
        workbook[sheet_name].sheet_state = "hidden"
    for sheet_name in VISIBLE_SHEETS:
        workbook[sheet_name].sheet_state = "visible"

    workbook.save(workbook_path)


def _theme_colors() -> dict[str, str]:
    return {
        "navy": "16324F",
        "teal": "2A7F62",
        "gold": "B88746",
        "mist": "EEF3F7",
        "line": "D7E1E8",
        "text": "182531",
        "warning": "C66A1A",
    }


def _apply_title_style(cell: Any) -> None:
    colors = _theme_colors()
    cell.font = Font(name="Aptos Display", size=18, bold=True, color=colors["navy"])


def _apply_header_style(cell: Any) -> None:
    colors = _theme_colors()
    cell.fill = PatternFill("solid", fgColor=colors["navy"])
    cell.font = Font(name="Aptos", size=10, bold=True, color="FFFFFF")
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = Border(bottom=Side(style="thin", color=colors["line"]))


def _apply_kpi_style(top_left: Any, label: str, value: Any) -> None:
    colors = _theme_colors()
    top_left.value = label
    top_left.font = Font(name="Aptos", size=9, bold=True, color=colors["navy"])
    top_left.fill = PatternFill("solid", fgColor=colors["mist"])
    top_left.border = Border(
        left=Side(style="thin", color=colors["line"]),
        right=Side(style="thin", color=colors["line"]),
        top=Side(style="thin", color=colors["line"]),
        bottom=Side(style="thin", color=colors["line"]),
    )
    value_cell = top_left.offset(row=1, column=0)
    value_cell.value = value
    value_cell.font = Font(name="Aptos Display", size=16, bold=True, color=colors["teal"])
    value_cell.fill = PatternFill("solid", fgColor="FFFFFF")
    value_cell.border = copy(top_left.border)


def _write_table(
    sheet: Any,
    start_row: int,
    start_col: int,
    title: str,
    rows: list[dict[str, Any]],
    columns: list[str],
    number_formats: dict[str, str] | None = None,
) -> tuple[int, int]:
    if title:
        title_cell = sheet.cell(row=start_row, column=start_col, value=title)
        title_cell.font = Font(name="Aptos", size=11, bold=True, color=_theme_colors()["navy"])
        start_row += 1
    for offset, column_name in enumerate(columns):
        cell = sheet.cell(row=start_row, column=start_col + offset, value=column_name)
        _apply_header_style(cell)
    data_start_row = start_row + 1
    for row_index, row in enumerate(rows, start=data_start_row):
        for col_index, column_name in enumerate(columns, start=start_col):
            value = row.get(column_name, "")
            cell = sheet.cell(row=row_index, column=col_index, value=value)
            cell.alignment = Alignment(vertical="top")
            cell.border = Border(bottom=Side(style="thin", color=_theme_colors()["line"]))
            if number_formats and column_name in number_formats:
                cell.number_format = number_formats[column_name]
    end_row = max(start_row, data_start_row + len(rows) - 1)
    for col_index, column_name in enumerate(columns, start=start_col):
        width = min(28, max(12, len(column_name) + 2))
        sheet.column_dimensions[_column_letter(col_index)].width = width
    if rows:
        sheet.auto_filter.ref = f"{_column_letter(start_col)}{start_row}:{_column_letter(start_col + len(columns) - 1)}{end_row}"
    return start_row, end_row


def _column_letter(index: int) -> str:
    result = ""
    while index:
        index, remainder = divmod(index - 1, 26)
        result = chr(65 + remainder) + result
    return result


def _populate_overview_sheet(
    sheet: Any,
    cluster_rows: list[dict[str, Any]],
    metric_rows: list[dict[str, Any]],
    config: DashboardBuildConfig,
    snapshot_date: str,
    marketplace: str,
    taxonomy_version: str,
) -> None:
    sheet["A1"] = "Storage Market Keyword Opportunity Dashboard"
    _apply_title_style(sheet["A1"])
    sheet["A2"] = f"Snapshot {snapshot_date} | Marketplace {marketplace} | Taxonomy {taxonomy_version}"
    sheet["A2"].font = Font(name="Aptos", size=10, color=_theme_colors()["text"])

    metric_lookup = {row["metric_name"]: row["metric_value"] for row in metric_rows}
    kpi_cells = ["A4", "C4", "E4", "G4", "I4"]
    kpis = [
        ("Total Keywords", metric_lookup.get("total_keywords", 0)),
        ("Ready Clusters", metric_lookup.get("ready_clusters", 0)),
        ("Candidate Clusters", metric_lookup.get("candidate_clusters", 0)),
        ("Watchlist Clusters", metric_lookup.get("watchlist_clusters", 0)),
        ("Main Pool Rows", metric_lookup.get("main_pool_keywords", 0)),
    ]
    for cell_name, (label, value) in zip(kpi_cells, kpis):
        _apply_kpi_style(sheet[cell_name], label, value)

    ready_top10 = [
        row
        for row in sorted(cluster_rows, key=lambda row: (-float(row["priority_score"]), -int(row["keyword_count"])))
        if row["include_in_top10"]
    ]
    _write_table(
        sheet,
        start_row=8,
        start_col=1,
        title="Top 10 Ready Clusters",
        rows=ready_top10,
        columns=[
            "ready_rank",
            "cluster_label",
            "priority_score",
            "keyword_count",
            "top100k_penetration",
            "best_rank",
            "representative_keyword",
        ],
        number_formats={"priority_score": "0.00", "top100k_penetration": "0.0%"},
    )

    scene_distribution = _distribution_rows(cluster_rows, "scene_label")
    need_distribution = _distribution_rows(cluster_rows, "function_label")
    _write_table(
        sheet,
        start_row=8,
        start_col=10,
        title="Scene Distribution",
        rows=scene_distribution,
        columns=["label", "cluster_count", "keyword_count"],
    )
    _write_table(
        sheet,
        start_row=8,
        start_col=14,
        title="Need Distribution",
        rows=need_distribution,
        columns=["label", "cluster_count", "keyword_count"],
    )
    sheet.freeze_panes = "A8"


def _distribution_rows(cluster_rows: list[dict[str, Any]], key_name: str) -> list[dict[str, Any]]:
    grouped: dict[str, dict[str, Any]] = {}
    for row in cluster_rows:
        label = str(row[key_name])
        bucket = grouped.setdefault(label, {"label": label, "cluster_count": 0, "keyword_count": 0})
        bucket["cluster_count"] += 1
        bucket["keyword_count"] += int(row["keyword_count"])
    return sorted(grouped.values(), key=lambda row: (-row["keyword_count"], row["label"]))[:12]


def _populate_matrix_sheet(sheet: Any, cluster_rows: list[dict[str, Any]]) -> None:
    sheet["A1"] = "Opportunity Cluster Matrix"
    _apply_title_style(sheet["A1"])
    sheet["A2"] = "Matrix value = best priority score in each Scene x Need bucket. Candidate clusters are visible here but excluded from Top 10."
    sheet["A2"].font = Font(name="Aptos", size=10, color=_theme_colors()["text"])

    scoped_rows = [row for row in cluster_rows if row["cluster_quality"] in {"ready", "candidate", "backlog"}]
    scenes = sorted({row["scene_label"] for row in scoped_rows})
    needs = sorted({row["function_label"] for row in scoped_rows})
    row_start = 5
    col_start = 2

    for col_offset, need_label in enumerate(needs, start=col_start):
        cell = sheet.cell(row=row_start, column=col_offset, value=need_label)
        _apply_header_style(cell)
    for row_offset, scene_label in enumerate(scenes, start=row_start + 1):
        cell = sheet.cell(row=row_offset, column=1, value=scene_label)
        _apply_header_style(cell)

    matrix_min_row = row_start + 1
    matrix_max_row = row_start + len(scenes)
    matrix_min_col = col_start
    matrix_max_col = col_start + len(needs) - 1

    matrix_values: dict[tuple[str, str], float] = {}
    for row in scoped_rows:
        key = (row["scene_label"], row["function_label"])
        matrix_values[key] = max(matrix_values.get(key, 0.0), float(row["priority_score"]))

    for row_index, scene_label in enumerate(scenes, start=matrix_min_row):
        for col_index, need_label in enumerate(needs, start=matrix_min_col):
            cell = sheet.cell(row=row_index, column=col_index, value=matrix_values.get((scene_label, need_label), 0.0))
            cell.number_format = "0.00"
            cell.alignment = Alignment(horizontal="center")

    if scenes and needs and ColorScaleRule is not None:
        matrix_ref = f"{_column_letter(matrix_min_col)}{matrix_min_row}:{_column_letter(matrix_max_col)}{matrix_max_row}"
        sheet.conditional_formatting.add(
            matrix_ref,
            ColorScaleRule(
                start_type="min",
                start_color="EEF3F7",
                mid_type="percentile",
                mid_value=50,
                mid_color="9BC6B8",
                end_type="max",
                end_color="2A7F62",
            ),
        )

    _write_table(
        sheet,
        start_row=max(18, matrix_max_row + 3),
        start_col=1,
        title="Opportunity Cluster Table",
        rows=sorted(scoped_rows, key=lambda row: (-float(row["priority_score"]), row["cluster_label"])),
        columns=[
            "cluster_label",
            "cluster_quality",
            "priority_score",
            "keyword_count",
            "top100k_penetration",
            "best_rank",
            "representative_keyword",
        ],
        number_formats={"priority_score": "0.00", "top100k_penetration": "0.0%"},
    )
    sheet.freeze_panes = "B6"


def _populate_drilldown_sheet(
    sheet: Any,
    bridge_rows: list[dict[str, Any]],
    cluster_rows: list[dict[str, Any]],
) -> None:
    sheet["A1"] = "Cluster Drilldown"
    _apply_title_style(sheet["A1"])
    sheet["A2"] = "Keyword-level evidence bridge. Use filters on cluster_label, row_outcome, and rank."
    sheet["A2"].font = Font(name="Aptos", size=10, color=_theme_colors()["text"])
    cluster_lookup = {row["cluster_key"]: row for row in cluster_rows}
    joined_rows = []
    for bridge_row in bridge_rows:
        cluster_row = cluster_lookup[bridge_row["cluster_key"]]
        joined_rows.append(
            {
                "cluster_label": bridge_row["cluster_label"],
                "cluster_quality": cluster_row["cluster_quality"],
                "priority_score": cluster_row["priority_score"],
                "keyword": bridge_row["keyword"],
                "rank": bridge_row["rank"],
                "row_outcome": bridge_row["row_outcome"],
                "scene_key": bridge_row["scene_key"],
                "function_key": bridge_row["function_key"],
                "form_rollup_key": bridge_row["form_rollup_key"],
                "matched_aliases": bridge_row["matched_aliases"],
                "evidence": bridge_row["evidence"],
                "unmapped_phrases": bridge_row["unmapped_phrases"],
            }
        )
    _write_table(
        sheet,
        start_row=5,
        start_col=1,
        title="Keyword Bridge",
        rows=joined_rows,
        columns=[
            "cluster_label",
            "cluster_quality",
            "priority_score",
            "keyword",
            "rank",
            "row_outcome",
            "scene_key",
            "function_key",
            "form_rollup_key",
            "matched_aliases",
            "evidence",
            "unmapped_phrases",
        ],
        number_formats={"priority_score": "0.00"},
    )
    sheet.freeze_panes = "A6"


def _populate_methodology_sheet(
    sheet: Any,
    config: DashboardBuildConfig,
    source_path: Path,
    snapshot_date: str,
    marketplace: str,
    taxonomy_version: str,
) -> None:
    sheet["A1"] = "Methodology And Definitions"
    _apply_title_style(sheet["A1"])
    rows = [
        ("Data source", str(source_path)),
        ("Decision unit", "Opportunity cluster"),
        ("Cluster key", "{cluster_key_version, anchor_type, anchor_value, object_key, form_rollup_key, function_key}"),
        ("Main pool rule", "mapped + partial rows"),
        ("Watchlist rule", "ambiguous rows or high-ambiguity clusters"),
        ("Top 10 rule", "Only ready clusters can enter Top 10"),
        ("Demand proxy", "keyword_count + Top100k penetration + sum(1/rank)"),
        ("Search volume treatment", "Used only when populated. Missing values do not suppress dashboard build."),
        ("V1 exclusions", "No supply-chain combos. No composite scoring block."),
        ("Visible sheets", ", ".join(VISIBLE_SHEETS)),
        ("Hidden contract sheets", ", ".join(HIDDEN_SHEETS)),
        ("Approach", "Excel-first dashboard with contract-compatible outputs"),
        ("Snapshot date", snapshot_date),
        ("Marketplace", marketplace),
        ("Taxonomy version", taxonomy_version),
    ]
    for row_index, (label, value) in enumerate(rows, start=4):
        label_cell = sheet.cell(row=row_index, column=1, value=label)
        value_cell = sheet.cell(row=row_index, column=2, value=value)
        label_cell.font = Font(name="Aptos", size=10, bold=True, color=_theme_colors()["navy"])
        value_cell.alignment = Alignment(wrap_text=True, vertical="top")
    sheet.column_dimensions["A"].width = 24
    sheet.column_dimensions["B"].width = 80


def _populate_contract_sheet(sheet: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        return
    columns = list(rows[0].keys())
    _write_table(sheet, start_row=1, start_col=1, title="", rows=rows, columns=columns)
    sheet.freeze_panes = "A2"


def _resolve_metadata_value(rows: list[dict[str, Any]], field_name: str, default: str) -> str:
    values = {_clean_text(row.get(field_name)) for row in rows if _clean_text(row.get(field_name))}
    if not values:
        return default
    if len(values) == 1:
        return values.pop()
    return f"mixed:{field_name}"


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Build the storage keyword opportunity dashboard from kw.csv.")
    parser.add_argument("--kw-csv", required=True, help="Path to kw.csv")
    parser.add_argument("--output-dir", help="Output directory for contract CSVs and workbook")
    parser.add_argument("--template-path", help="Optional workbook template or style reference")
    parser.add_argument("--snapshot-date", help="Snapshot date written into artifacts")
    parser.add_argument("--marketplace", default="US")
    parser.add_argument("--taxonomy-version", default="storage_v1")
    parser.add_argument("--cluster-key-version", default="v1")
    parser.add_argument("--workbook-name", default=DEFAULT_WORKBOOK_NAME)
    parser.add_argument("--no-workbook", action="store_true", help="Write contract CSVs only")
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    kw_csv_path = Path(args.kw_csv).expanduser().resolve()
    output_dir = Path(args.output_dir).expanduser().resolve() if args.output_dir else kw_csv_path.parent / "dashboard_v1"
    template_path = Path(args.template_path).expanduser().resolve() if args.template_path else None
    config = DashboardBuildConfig(
        kw_csv_path=kw_csv_path,
        output_dir=output_dir,
        template_path=template_path,
        snapshot_date=args.snapshot_date,
        marketplace=args.marketplace,
        taxonomy_version=args.taxonomy_version,
        cluster_key_version=args.cluster_key_version,
        workbook_name=args.workbook_name,
        write_workbook=not args.no_workbook,
    )
    artifacts = build_keyword_dashboard(config)
    printable = {name: str(path) for name, path in artifacts.items() if config.write_workbook or name != "workbook"}
    print(json.dumps(printable, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
